"""documente_doc — un model minimal de document (titluri, paragrafe, câmpuri, tabele) randat identic în .docx
(python-docx) și HTML (→ PDF prin Chrome). Folosit de dosarul de racordare și de caietul de sarcini (E2).

Convenția câmpurilor (cerința proprietarului — „marchează clar câmpurile completate din date vs de completat manual”):
* Camp(valoare)                    completat AUTOMAT din datele proiectului → fond albastru deschis;
* Camp(None) / Camp("", manual)    DE COMPLETAT MANUAL → „……… (de completat)” pe fond galben.
Fiecare document începe cu legenda celor două marcaje.
"""
import io
import re

CULOARE_DATE = "DCEBFA"        # albastru deschis
CULOARE_MANUAL = "FFF2A8"      # galben


class Camp:
    """o valoare dintr-un formular: completată din date sau de completat manual"""

    def __init__(self, valoare=None, nota=None):
        v = "" if valoare is None else str(valoare).strip()
        self.valoare = v
        self.manual = not v
        self.nota = nota

    def text(self):
        return self.valoare if not self.manual else "…………………… (de completat)"


def camp(v, unitate=""):
    if v is None or v == "" or v is False:
        return Camp(None)
    if isinstance(v, float):
        v = (f"{v:.2f}".rstrip("0").rstrip(".")).replace(".", ",")
    return Camp(f"{v}{(' ' + unitate) if unitate else ''}")


class Document:
    def __init__(self, titlu, subtitlu="", antet=""):
        self.titlu, self.subtitlu, self.antet = titlu, subtitlu, antet
        self.blocuri = []

    # blocuri ------------------------------------------------------------------------------------------------
    def h(self, text, nivel=1):
        self.blocuri.append(("h", nivel, text))
        return self

    def p(self, *bucati, stil=None):
        """paragraf din bucăți: șiruri (text fix) și Camp-uri"""
        self.blocuri.append(("p", stil, list(bucati)))
        return self

    def lista(self, elemente, numerotata=False):
        self.blocuri.append(("l", numerotata, list(elemente)))
        return self

    def campuri(self, perechi):
        """formular pe două coloane: [(eticheta, Camp | text), …]"""
        self.blocuri.append(("f", None, list(perechi)))
        return self

    def tabel(self, antet, randuri, latimi=None):
        self.blocuri.append(("t", latimi, (list(antet), [list(r) for r in randuri])))
        return self

    def semnaturi(self, *roluri):
        self.blocuri.append(("s", None, list(roluri)))
        return self

    def salt(self):
        self.blocuri.append(("br", None, None))
        return self

    def legenda(self):
        self.blocuri.append(("leg", None, None))
        return self

    # statistică ---------------------------------------------------------------------------------------------
    def numar_campuri(self):
        auto = manual = 0

        def num(x):
            nonlocal auto, manual
            if isinstance(x, Camp):
                if x.manual:
                    manual += 1
                else:
                    auto += 1
            elif isinstance(x, (list, tuple)):
                for y in x:
                    num(y)
        for b in self.blocuri:
            num(b[2])
        return {"completate": auto, "de_completat": manual}

    # HTML ---------------------------------------------------------------------------------------------------
    def html(self):
        def esc(s):
            return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        def bucata(x):
            if isinstance(x, Camp):
                cls = "m" if x.manual else "d"
                return f"<span class='{cls}'>{esc(x.text())}</span>"
            return esc(x)
        css = ("@page{size:A4;margin:18mm 16mm}body{font:10.5pt/1.45 'Barlow','Arial',sans-serif;color:#111}"
               "h1{font-size:15pt;text-align:center;margin:0 0 1mm}h2{font-size:12pt;margin:6mm 0 2mm;border-bottom:.3mm solid #999}"
               "h3{font-size:11pt;margin:4mm 0 1.5mm}.sub{text-align:center;color:#444;margin:0 0 5mm}.antet{color:#555;font-size:9pt;"
               "border-bottom:.3mm solid #bbb;margin-bottom:4mm;padding-bottom:1mm}"
               f".d{{background:#{CULOARE_DATE};padding:0 .6mm}}.m{{background:#{CULOARE_MANUAL};padding:0 .6mm;color:#6a5600}}"
               "table{border-collapse:collapse;width:100%;margin:2mm 0 3mm}td,th{border:.3mm solid #888;padding:1mm 1.5mm;"
               "vertical-align:top;font-size:9.5pt}th{background:#eeeeea;text-align:left}table.f td:first-child{width:42%;"
               "color:#333}.leg{font-size:8.5pt;border:.3mm solid #ccc;padding:1.5mm 2mm;margin:0 0 4mm}.sig{display:flex;"
               "justify-content:space-between;gap:8mm;margin-top:12mm}.sig div{flex:1;border-top:.3mm solid #555;padding-top:1mm;"
               "font-size:9pt;text-align:center}.br{page-break-after:always}p{margin:1.5mm 0}")
        h = [f"<!doctype html><html lang='ro'><head><meta charset='utf-8'><title>{esc(self.titlu)}</title>"
             f"<style>{css}</style></head><body>"]
        if self.antet:
            h.append(f"<div class='antet'>{esc(self.antet)}</div>")
        h.append(f"<h1>{esc(self.titlu)}</h1>")
        if self.subtitlu:
            h.append(f"<p class='sub'>{esc(self.subtitlu)}</p>")
        for tip, a, b in self.blocuri:
            if tip == "h":
                h.append(f"<h{a + 1}>{esc(b)}</h{a + 1}>")
            elif tip == "p":
                h.append("<p>" + "".join(bucata(x) for x in b) + "</p>")
            elif tip == "l":
                tag = "ol" if a else "ul"
                h.append(f"<{tag}>" + "".join("<li>" + ("".join(bucata(y) for y in x) if isinstance(x, (list, tuple))
                                                        else bucata(x)) + "</li>" for x in b) + f"</{tag}>")
            elif tip == "f":
                h.append("<table class='f'>" + "".join(f"<tr><td>{esc(e)}</td><td>{bucata(v)}</td></tr>" for e, v in b) + "</table>")
            elif tip == "t":
                ant, rows = b
                h.append("<table><thead><tr>" + "".join(f"<th>{esc(x)}</th>" for x in ant) + "</tr></thead><tbody>"
                         + "".join("<tr>" + "".join(f"<td>{bucata(x)}</td>" for x in r) + "</tr>" for r in rows)
                         + "</tbody></table>")
            elif tip == "s":
                h.append("<div class='sig'>" + "".join(f"<div>{esc(x)}<br>(nume, semnătura, ștampila)</div>" for x in b) + "</div>")
            elif tip == "br":
                h.append("<div class='br'></div>")
            elif tip == "leg":
                h.append(f"<div class='leg'><span class='d'>Text pe fond albastru</span> = completat automat din datele "
                         f"proiectului (de verificat). <span class='m'>…… (de completat)</span> = câmp de completat manual "
                         f"înainte de depunere / semnare.</div>")
        h.append("</body></html>")
        return "".join(h)

    # DOCX ---------------------------------------------------------------------------------------------------
    def docx(self):
        """bytes .docx (python-docx); ImportError dacă biblioteca lipsește (versiunea web → se folosește HTML)"""
        from docx import Document as D
        from docx.shared import Pt, Mm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        d = D()
        sec = d.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(18)
        sec.top_margin = sec.bottom_margin = Mm(16)
        st = d.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(10.5)

        def umbra(run, culoare):
            rpr = run._r.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), culoare)
            rpr.append(shd)

        def scrie(par, x, bold=False):
            if isinstance(x, Camp):
                r = par.add_run(x.text())
                umbra(r, CULOARE_MANUAL if x.manual else CULOARE_DATE)
                if x.manual:
                    r.font.color.rgb = RGBColor(0x6A, 0x56, 0x00)
            else:
                r = par.add_run(str(x))
            r.bold = bold
            return r

        def celula_umbra(cel, culoare):
            tcpr = cel._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), culoare)
            tcpr.append(shd)
        if self.antet:
            pa = sec.header.paragraphs[0]
            pa.text = self.antet
            pa.runs[0].font.size = Pt(8)
        t = d.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run(self.titlu)
        r.bold = True
        r.font.size = Pt(15)
        if self.subtitlu:
            s = d.add_paragraph()
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.add_run(self.subtitlu).italic = True
        for tip, a, b in self.blocuri:
            if tip == "h":
                d.add_heading(b, level=min(a, 3))
            elif tip == "p":
                par = d.add_paragraph()
                for x in b:
                    scrie(par, x)
            elif tip == "l":
                for x in b:
                    par = d.add_paragraph(style="List Number" if a else "List Bullet")
                    for y in (x if isinstance(x, (list, tuple)) else [x]):
                        scrie(par, y)
            elif tip in ("f", "t"):
                if tip == "f":
                    ant, rows = None, [[e, v] for e, v in b]
                else:
                    ant, rows = b
                nc = len(ant or rows[0]) if (ant or rows) else 1
                tb = d.add_table(rows=0, cols=nc)
                tb.style = "Table Grid"
                if ant:
                    cells = tb.add_row().cells
                    for i, x in enumerate(ant):
                        cells[i].text = ""
                        scrie(cells[i].paragraphs[0], x, bold=True)
                        celula_umbra(cells[i], "EEEEEA")
                for rw in rows:
                    cells = tb.add_row().cells
                    for i, x in enumerate(rw[:nc]):
                        cells[i].text = ""
                        scrie(cells[i].paragraphs[0], x)
                for row in tb.rows:
                    for c in row.cells:
                        for pp in c.paragraphs:
                            for rr in pp.runs:
                                rr.font.size = Pt(9)
                d.add_paragraph()
            elif tip == "s":
                tb = d.add_table(rows=2, cols=len(b))
                for i, x in enumerate(b):
                    tb.cell(0, i).text = x
                    tb.cell(1, i).text = "\n\n(nume, semnătura, ștampila)"
            elif tip == "br":
                d.add_page_break()
            elif tip == "leg":
                par = d.add_paragraph()
                r1 = par.add_run("Text pe fond albastru")
                umbra(r1, CULOARE_DATE)
                par.add_run(" = completat automat din datele proiectului (de verificat).  ")
                r2 = par.add_run("…… (de completat)")
                umbra(r2, CULOARE_MANUAL)
                par.add_run(" = câmp de completat manual înainte de depunere / semnare.")
                for rr in par.runs:
                    rr.font.size = Pt(8.5)
        bio = io.BytesIO()
        d.save(bio)
        return bio.getvalue()


def docx_disponibil():
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def nume_sigur(t):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", t).strip("_")
