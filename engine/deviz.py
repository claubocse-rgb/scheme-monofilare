"""Deviz estimativ generat din lista de echipamente și materiale a planșei.

* Articolele = pozițiile listei (aceleași ca pe planșă și în BOM): denumire, caracteristici, cantitate, U.M.
* Prețurile vin dintr-un CATALOG local (`date/preturi.json`, lângă aplicație), indexat după denumirea articolului.
  Un preț introdus o dată se regăsește automat în toate proiectele următoare care folosesc același articol.
* Manopera: ore normate pe unitate (implicit pe categorii — valori ORIENTATIVE, editabile) × tariful orar.
* Recapitulația: materiale + manoperă → cheltuieli indirecte % → profit % → total fără TVA → TVA → total.
Nimic nu e inventat: un articol fără preț în catalog rămâne fără preț (marcat), iar totalul spune câte lipsesc.
"""
import datetime, io, json, os, pathlib, re

ROOT = pathlib.Path(__file__).resolve().parents[1]
DATE = pathlib.Path(os.environ.get("SM_DATE") or ROOT / "date")     # datele utilizatorului (catalog, profile) — în afara git
CATALOG = DATE / "preturi.json"

SETARI_IMPLICITE = {"moneda": "lei", "tarif_ora": 0.0, "ch_indirecte_pct": 0.0, "profit_pct": 0.0, "tva_pct": 21.0,
                    "aplica_ore_implicite": True}

# ore de manoperă pe unitate, pe categorie — ORIENTATIVE (montaj + conectare + probe), se suprascriu din catalog / UI
ORE_IMPLICITE = [
    (r"^cablu .*(\d+)×(\d+)", "cablu_m"),
    (r"^conductor de protecție", 0.03),
    (r"^set accesorii capăt", 0.4),
    (r"^patron fuzibil", 0.05),
    (r"^(întreruptor automat de putere|acb)", 6.0),
    (r"^întreruptor automat compact", 2.0),
    (r"^întreruptor automat (modular|mcb)|^întreruptor automat .*mcb|^mcb", 0.4),
    (r"^întreruptor (automat )?diferențial|^rccb|^bloc diferențial|^întreruptor automat cu protecție diferențială", 0.6),
    (r"^separator", 1.0),
    (r"^întreruptor-separator", 0.4),
    (r"^tablou modular", 4.0),
    (r"^pieptene", 0.3),
    (r"^descărcător", 0.6),
    (r"^contor", 1.0),
    (r"^transformator de curent", 0.5),
    (r"^șir de cleme|^bloc siguranțe", 0.5),
    (r"^sistem de bare", 4.0),
    (r"^carcasă tablou", 6.0),
    (r"^invertor", 4.0),
    (r"^comutator automat de sursă", 6.0),
    (r"^electrod de pământ", 1.0),
    (r"^platbandă", 0.15),
    (r"^clemă|^îmbinare", 0.3),
    (r"^săpătură", 0.6),
    (r"^vopsea", 0.0),
]


def cheie(denumire):
    """cheia de catalog: denumirea normalizată (fără diferențe de spații / majuscule / „x” vs „×” / „mm2” vs „mm²”)"""
    t = str(denumire or "").replace("\u00a0", " ").replace("\u202f", " ").lower()
    t = re.sub(r"(?<=\d)\s*[x×]\s*(?=\d)", "×", t)
    t = re.sub(r"\bmm\s*(2|p|²)(?=\W|$)", "mm²", t)
    return " ".join(t.split())


def cheie_lax(denumire):
    """potrivire tolerantă: fără material la final (Al/Cu), fără punctuație și fără „mm²”"""
    t = cheie(denumire)
    t = re.sub(r"\s+(al|cu)$", "", t)
    t = t.replace("mm²", "")
    t = re.sub(r"[^\w×+]+", " ", t)
    return " ".join(t.split())


def _ore_implicite(den, um):
    k = cheie(den)
    for pat, v in ORE_IMPLICITE:
        m = re.search(pat, k)
        if not m:
            continue
        if v == "cablu_m":           # cablu: crește cu secțiunea (pozare + tragere), pe metru
            try:
                S = float(m.group(2))
            except (TypeError, ValueError):
                S = 16
            return 0.05 if S <= 16 else 0.1 if S <= 70 else 0.15 if S <= 150 else 0.2
        return v
    return None


def citeste_catalog():
    try:
        d = json.loads(CATALOG.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        d = {}
    d.setdefault("articole", {})
    d["setari"] = dict(SETARI_IMPLICITE, **(d.get("setari") or {}))
    return d


def scrie_catalog(nou):
    """îmbină în catalog articolele și setările primite ({articole: {cheie: {...}}, setari: {...}})"""
    d = citeste_catalog()
    azi = datetime.date.today().isoformat()
    for k, v in (nou.get("articole") or {}).items():
        k = cheie(k)
        if not k:
            continue
        if v is None:
            d["articole"].pop(k, None); continue
        e = dict(d["articole"].get(k) or {})
        for kk in ("pret", "ore"):
            if kk in v and v[kk] is not None:
                n = _num(v[kk])
                if n is None or n < 0:
                    raise ValueError(f"{kk} „{v[kk]}” pentru „{v.get('denumire') or k}” nu e un număr pozitiv")
                e[kk] = n
        for kk in ("furnizor", "cod", "denumire"):
            if kk in v:
                e[kk] = str(v[kk] or "")
        e["actualizat"] = azi
        d["articole"][k] = e
    if nou.get("setari"):
        d["setari"].update(_valideaza_setari(nou["setari"]))
    DATE.mkdir(exist_ok=True)
    CATALOG.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
    return d


def _num(x):
    """număr din celulă sau text în format românesc: 1.250,50 · 1.250 · 111,5 lei · 12.5 · 1 250"""
    if x is None or isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return float(x) if x == x else None
    t = re.sub(r"(?i)\s*(lei|ron|eur|€|/\s*\w+)\s*$", "", str(x).strip()).replace("\u00a0", "").replace(" ", "")
    if not t:
        return None
    if "," in t and "." in t:              # separatorul zecimal e ultimul: 1.250,50 (RO) sau 1,250.50 (EN)
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3}){2,}", t):   # 1,250,000
        t = t.replace(",", "")
    elif "," in t:                          # 111,5
        t = t.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", t):   # 1.250 = o mie două sute cincizeci
        t = t.replace(".", "")
    try:
        v = float(t)
        return v if v == v else None
    except ValueError:
        return None


def _valideaza_setari(st):
    """setările devizului: numere finite, procente în [0, 100], tarif ≥ 0; altfel ValueError cu mesaj clar"""
    out = {}
    lim = {"tarif_ora": (0, 1e6), "ch_indirecte_pct": (0, 100), "profit_pct": (0, 100), "tva_pct": (0, 100)}
    for k, v in (st or {}).items():
        if k not in SETARI_IMPLICITE:
            continue
        if k in lim:
            n = _num(v)
            if n is None or not (lim[k][0] <= n <= lim[k][1]):
                raise ValueError(f"valoarea „{v}” pentru {k} nu e acceptată (interval {lim[k][0]}…{lim[k][1]})")
            out[k] = n
        elif k == "aplica_ore_implicite":
            out[k] = bool(v)
        elif k == "moneda":
            out[k] = str(v)[:8] or "lei"
    return out


def articole(poz):
    """pozițiile listei (nr, ref, lista, buc, tip) → articole de deviz"""
    out = []
    for (n, ref, lista, buc, tip) in poz:
        den = lista[0] if lista else ""
        car = " · ".join(lista[1:]) if len(lista) > 1 else ""
        um = "buc"
        if isinstance(buc, str) and " " in buc:
            q, um = buc.split(" ", 1)
            buc = _num(q)
        out.append({"poz": n, "ref": ref, "denumire": den, "caracteristici": car, "cantitate": _num(buc) or 0, "um": um,
                    "tip": tip, "cheie": cheie(den)})
    return out


def calculeaza(poz, catalog=None, setari=None, suprascrieri=None, extra=None):
    """articole cu prețuri și ore + recapitulație.
    `suprascrieri` = {cheie: {pret, ore, cantitate, exclus}} din interfață (nesalvate în catalog);
    `extra` = articole adăugate manual [{denumire, um, cantitate, pret, ore, caracteristici}] (structură, transport, PIF …)."""
    cat = catalog or citeste_catalog()
    st = dict(SETARI_IMPLICITE)
    try:
        st.update(_valideaza_setari(cat.get("setari")))
    except ValueError:
        pass                                   # catalog vechi cu valori greșite → implicitele
    st.update(_valideaza_setari(setari))       # setările proiectului: greșite → eroare clară
    sup = {cheie(k): v for k, v in (suprascrieri or {}).items()}
    lax = {}
    for k, v in cat["articole"].items():
        lax.setdefault(cheie_lax(k), v)
    rows, mat, ore_tot, lipsa = [], 0.0, 0.0, 0
    art = articole(poz)
    for i, x in enumerate(extra or []):
        if not str(x.get("denumire") or "").strip():
            continue
        q = _num(x.get("cantitate"))
        if q is None or q < 0:
            raise ValueError(f"cantitatea articolului adăugat „{x.get('denumire')}” nu e un număr pozitiv")
        art.append({"poz": f"+{i + 1}", "ref": "manual", "denumire": str(x["denumire"]).strip(), "caracteristici": str(x.get("caracteristici") or ""),
                    "cantitate": q, "um": str(x.get("um") or "buc"), "tip": "extra", "cheie": cheie(x["denumire"]), "extra_idx": i,
                    "_pret": x.get("pret"), "_ore": x.get("ore")})
    for a in art:
        s_ = sup.get(a["cheie"]) or {}
        e = dict(cat["articole"].get(a["cheie"]) or lax.get(cheie_lax(a["denumire"])) or {})
        if a.get("tip") == "extra":
            pe, oe = a.pop("_pret"), a.pop("_ore")
            if pe not in (None, ""): e["pret"] = pe
            if oe not in (None, ""): e["ore"] = oe
        a["cantitate_plansa"] = a["cantitate"]
        if s_.get("cantitate") not in (None, "") and a.get("tip") != "extra":
            q = _num(s_["cantitate"])
            if q is None or q < 0:
                raise ValueError(f"cantitatea pentru „{a['denumire']}” nu e un număr pozitiv")
            a["cantitate"] = q
        a["exclus"] = bool(s_.get("exclus"))
        e.update({k: v for k, v in s_.items() if k in ("pret", "ore") and v not in (None, "")})
        pret = _num(e.get("pret"))
        ore = _num(e.get("ore"))
        if pret is not None and pret < 0:
            raise ValueError(f"prețul pentru „{a['denumire']}” este negativ")
        if ore is not None and ore < 0:
            raise ValueError(f"orele pentru „{a['denumire']}” sunt negative")
        ore_sursa = "catalog" if ore is not None else None
        if ore is None and st.get("aplica_ore_implicite"):
            ore = _ore_implicite(a["denumire"], a["um"])
            ore_sursa = "orientativ" if ore is not None else None
        val = round(pret * a["cantitate"], 2) if pret is not None else None
        o = round((ore or 0) * a["cantitate"], 2)
        if a["exclus"]:                         # ex. contorul furnizat de OD: rămâne pe listă, nu intră în deviz
            val, o = (0.0 if pret is not None else None), 0.0
        elif pret is None:
            lipsa += 1
        else:
            mat += val
        ore_tot += o
        rows.append(dict(a, pret=pret, valoare=val, ore_u=ore, ore=o, ore_sursa=ore_sursa,
                         furnizor=e.get("furnizor", ""), cod=e.get("cod", ""), actualizat=e.get("actualizat", "")))
    man = round(ore_tot * float(st.get("tarif_ora") or 0), 2)
    baza = mat + man
    ind = round(baza * float(st.get("ch_indirecte_pct") or 0) / 100, 2)
    prof = round((baza + ind) * float(st.get("profit_pct") or 0) / 100, 2)
    fara = round(baza + ind + prof, 2)
    tva = round(fara * float(st.get("tva_pct") or 0) / 100, 2)
    fara_ore = sum(1 for r in rows if r["ore_u"] is None and not r["exclus"])
    rec = {"fara_ore": fara_ore, "materiale": round(mat, 2), "ore": round(ore_tot, 2), "manopera": man, "ch_indirecte": ind, "profit": prof,
           "total_fara_tva": fara, "tva": tva, "total": round(fara + tva, 2), "fara_pret": lipsa, "nr_articole": len(rows)}
    return {"articole": rows, "setari": st, "recapitulatie": rec}


def importa_xlsx(data: bytes):
    """listă de prețuri Excel (orice foaie): coloane recunoscute după antet — Denumire / Preț / Ore / Furnizor / Cod.
    Întoarce {articole: {cheie: {...}}} gata de îmbinat în catalog."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out, sarite = {}, []
    for ws in wb.worksheets:
        col = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(c or "").strip().lower() for c in row]
            if col is None:
                def idx(*nume, fara=()):
                    for i, c in enumerate(cells):
                        if any(c.startswith(n) or n in c for n in nume) and not any(f in c for f in fara):
                            return i
                    return None
                d = idx("denumire", "articol", "produs")
                # prețul UNITAR FĂRĂ TVA; coloanele „total”, „valoare”, „cu tva” nu se iau niciodată
                p = idx("preț net", "pret net", "preț unitar net", "pret unitar net", fara=("total", "valoare", "cu tva"))
                if p is None:
                    p = idx("preț unitar fără", "pret unitar fara", "preț unitar", "pret unitar", "pu ", "p.u.", fara=("total", "valoare", "cu tva", "listă", "lista"))
                if p is None:
                    p = idx("preț unitar", "pret unitar", "preț listă", "pret lista", fara=("total", "valoare", "cu tva"))
                if p is None:
                    p = idx("preț fără tva", "pret fara tva", "preț", "pret", fara=("total", "valoare", "cu tva"))
                if d is not None and p is None and any("cu tva" in c for c in cells):
                    raise ValueError("lista are doar prețuri cu TVA — adăugați o coloană „Preț unitar fără TVA”")
                if d is not None and p is not None:
                    col = {"den": d, "pret": p, "ore": idx("ore", "manoperă", "manopera", fara=("total",)),
                           "furnizor": idx("furnizor", "producător", "producator"), "cod": idx("cod")}
                continue
            den = row[col["den"]] if col["den"] < len(row) else None
            brut = row[col["pret"]] if col["pret"] < len(row) else None
            pret = _num(brut)
            if not den:
                continue
            if pret is None or pret < 0:
                if brut not in (None, ""):
                    sarite.append(f"{den}: „{brut}”")
                continue
            e = {"pret": pret, "denumire": str(den).strip()}
            for k in ("ore", "furnizor", "cod"):
                i = col.get(k)
                if i is not None and i < len(row) and row[i] not in (None, ""):
                    e[k] = _num(row[i]) if k == "ore" else str(row[i]).strip()
            out[cheie(den)] = e
    if not out:
        raise ValueError("nu am găsit rânduri cu denumire și preț unitar (antete recunoscute: Denumire / Articol, Preț unitar [fără TVA])")
    return {"articole": out, "sarite": sarite}


def deviz_xlsx(model, rezultat) -> bytes:
    """Excel cu formule (se poate continua în Excel): foaia „Deviz” + „Recapitulație”."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    k = model.get("cartus") or {}
    st = rezultat["setari"]; mon = st.get("moneda", "lei")
    wb = Workbook(); ws = wb.active; ws.title = "Deviz"
    thin = Side(style="thin", color="999999"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "DEVIZ ESTIMATIV — echipamente, materiale și manoperă"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{k.get('denumire2', '')}"
    ws["A3"] = f"Beneficiar: {k.get('beneficiar', '')} · Amplasament: {k.get('adresa', '')} · Proiect nr. {k.get('nr_pr', '')} · Planșa {k.get('plansa', '')} rev. {k.get('rev', '')}"
    ws["A4"] = f"Întocmit: {datetime.date.today().strftime('%d.%m.%Y')} · prețuri în {mon}, fără TVA · cantitățile cablurilor includ 5 % rezervă"
    hdr = ["Nr.", "Simbol", "Denumire", "Caracteristici", "U.M.", "Cant.", f"Preț unitar [{mon}]", f"Valoare materiale [{mon}]",
           "Ore / U.M.", "Ore total", "Furnizor / cod"]
    ws.append([]); ws.append(hdr)
    h = 6
    for c in ws[h]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="EEEEEA"); c.border = bd
        c.alignment = Alignment(vertical="center", wrap_text=True)
    galben = PatternFill("solid", fgColor="FFF2B3")
    r0 = h + 1
    for i, a in enumerate(rezultat["articole"]):
        r = r0 + i
        car = ("[EXCLUS din deviz — furnizat de alții] " if a.get("exclus") else "") + (a["caracteristici"] or "")
        ws.append([a["poz"], a["ref"], a["denumire"], car, a["um"], 0 if a.get("exclus") else a["cantitate"], a["pret"],
                   f"=IF(G{r}=\"\",\"\",F{r}*G{r})", a["ore_u"], f"=F{r}*N(I{r})",
                   " · ".join(x for x in (a.get("furnizor"), a.get("cod")) if x)])
        for c in ws[r]:
            c.border = bd; c.alignment = Alignment(vertical="top", wrap_text=True)
        if a["pret"] is None:
            ws[f"G{r}"].fill = galben
        if a.get("ore_sursa") == "orientativ":
            ws[f"I{r}"].font = Font(italic=True, color="7A6A00")
        for col in ("G", "H"):
            ws[f"{col}{r}"].number_format = '#,##0.00'
    rl = r0 + len(rezultat["articole"]) - 1
    for i, w in enumerate([5, 12, 52, 52, 6, 8, 13, 15, 9, 9, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r0}"
    n = rl + 2
    ws[f"C{n}"] = "Celulele galbene: articole fără preț în catalog — completați-le (totalul le ignoră până atunci)."
    ws[f"C{n + 1}"] = "Orele în cursiv sunt valori orientative pe categorie; înlocuiți-le cu normele proprii."
    for c in (f"C{n}", f"C{n + 1}"):
        ws[c].font = Font(italic=True, color="666666")

    rc = wb.create_sheet("Recapitulație")
    rc["A1"] = "RECAPITULAȚIE"; rc["A1"].font = Font(bold=True, size=13)
    rows = [
        ("Materiale și echipamente", f"=SUM(Deviz!H{r0}:H{rl})"),
        ("Ore manoperă", f"=SUM(Deviz!J{r0}:J{rl})"),
        (f"Tarif orar manoperă [{mon}/oră]", st.get("tarif_ora") or 0),
        ("Manoperă", "=B4*B5"),
        ("Total direct (materiale + manoperă)", "=B3+B6"),
        ("Cheltuieli indirecte [%]", st.get("ch_indirecte_pct") or 0),
        ("Cheltuieli indirecte", "=B7*B8/100"),
        ("Profit [%]", st.get("profit_pct") or 0),
        ("Profit", "=(B7+B9)*B10/100"),
        ("TOTAL fără TVA", "=B7+B9+B11"),
        ("TVA [%]", st.get("tva_pct") or 0),
        ("TVA", "=B12*B13/100"),
        ("TOTAL cu TVA", "=B12+B14"),
    ]
    for i, (a, b) in enumerate(rows, start=3):
        rc[f"A{i}"] = a; rc[f"B{i}"] = b
        rc[f"A{i}"].border = bd; rc[f"B{i}"].border = bd
        if not str(a).endswith("[%]") and "Ore" not in a and "Tarif" not in a:
            rc[f"B{i}"].number_format = '#,##0.00'
        if a.startswith("TOTAL"):
            rc[f"A{i}"].font = Font(bold=True); rc[f"B{i}"].font = Font(bold=True)
    for c in ("B5", "B8", "B10", "B13"):
        rc[c].fill = PatternFill("solid", fgColor="E8F0FF")
    rc["A18"] = f"Articole fără preț: {rezultat['recapitulatie']['fara_pret']} din {rezultat['recapitulatie']['nr_articole']}."
    rc["A19"] = "Celulele albastre se pot modifica; totalurile se recalculează."
    rc.column_dimensions["A"].width = 42; rc.column_dimensions["B"].width = 18
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
