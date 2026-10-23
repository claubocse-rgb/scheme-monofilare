"""Export listă de echipamente (BOM) în Excel, din pozițiile generate de layout."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def bom_xlsx(model, poz) -> bytes:
    """poz = lista (nr, ref, lista_lines, buc, tip) din Layout.poz"""
    wb = Workbook(); ws = wb.active; ws.title = "Lista echipamente"
    k = model.get("cartus", {})
    ws["A1"] = "LISTĂ DE ECHIPAMENTE ȘI MATERIALE"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{k.get('denumire1', '')} — {k.get('denumire2', '')}"; ws["A3"] = f"Beneficiar: {k.get('beneficiar', '')} · Proiect: {k.get('titlu', '')} · Planșa {k.get('plansa', '')} rev. {k.get('rev', '')}"
    hdr = ["Poz.", "Simbol", "Denumire", "Caracteristici tehnice", "Tip simbol", "Buc.", "U.M.", "Producător / cod", "Preț unitar", "Valoare"]
    ws.append([]); ws.append(hdr)
    thin = Side(style="thin", color="999999"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[5]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="EEEEEA"); c.border = border; c.alignment = Alignment(vertical="center", wrap_text=True)
    r = 6
    for (n, ref, lista, buc, tip) in poz:
        den = lista[0] if lista else ""; car = " · ".join(lista[1:]) if len(lista) > 1 else ""
        um = "buc"
        if isinstance(buc, str) and " " in buc:          # materiale: „105 m”
            q, um = buc.split(" ", 1); qq = q.replace(",", "."); buc = float(qq) if qq.replace(".", "", 1).isdigit() else q
        ws.append([n, ref, den, car, tip, buc, um, "", None, f"=F{r}*I{r}"])
        for c in ws[r]: c.border = border; c.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
    ws.append([]); ws.append(["", "", "TOTAL", "", "", "", "", "", "", f"=SUM(J6:J{r-1})"]); ws[f"J{r+1}"].font = Font(bold=True); ws[f"C{r+1}"].font = Font(bold=True)
    for i, w in enumerate([6, 12, 60, 70, 14, 6, 6, 24, 12, 14], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
