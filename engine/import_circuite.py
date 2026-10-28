"""import_circuite — modelul Excel descărcabil și importul circuitelor din Excel / CSV (E2, funcționalitatea #8).

    model_xlsx()                       → bytes: „Model import circuite.xlsx” (foaia Circuite + Receptori + Instrucțiuni)
    importa(date, destinatie=None)     → {circuite (locuință), plecari (tablou), avertismente, ignorate, n}

Parserul este tolerant: antetul se recunoaște după cuvinte (fără diacritice / majuscule, oriunde în primele 10 rânduri),
numerele pot avea virgulă („2,5”), unitate („16 A”, „3 kW”, „25m”) sau puterea în W („2500 W”); receptorul se recunoaște
după cheia internă sau după denumirea românească / sinonime; rândurile goale și cele de total se sar.
CSV: separatorul „;”, „,” sau tab, codarea UTF-8 (cu sau fără BOM) sau Windows-1250.
"""
import csv
import io
import re
import unicodedata

from .calc.locuinta import ETICHETE_RECEPTOR, FAZE

POZARI = {"B2": "în tub pe / în perete", "C": "aparent pe perete", "D1": "în tub în pământ", "D2": "direct în pământ",
          "E": "pe pat de cabluri / în aer"}
COLOANE = [("nume", "Circuit (denumire)", 28), ("receptor", "Tip receptor", 26), ("P_kW", "P [kW]", 9),
           ("cosfi", "cos φ", 7), ("L_m", "L [m]", 8), ("faza", "Fază (L1/L2/L3/3F)", 12), ("metoda", "Pozare (B2/C/D1/D2/E)", 14),
           ("In_A", "In impus [A]", 11), ("curba", "Curbă (B/C/D)", 10), ("obs", "Observații", 30)]
SINONIME_COLOANE = {
    "nume": ("circuit", "denumire", "nume", "consumator", "receptor denumire", "descriere", "plecare", "destinatie"),
    "receptor": ("tip receptor", "tip", "receptor", "categorie", "tip consumator"),
    "P_kW": ("p [kw]", "p(kw)", "p kw", "putere", "p instalat", "pi", "p", "kw", "putere instalata", "pn"),
    "cosfi": ("cos", "cosfi", "cos phi", "cos φ", "factor de putere", "pf"),
    "L_m": ("l [m]", "l(m)", "lungime", "l", "m", "distanta", "lungime cablu"),
    "faza": ("faza", "faze", "fază", "phase", "u", "retea"),
    "metoda": ("pozare", "metoda", "mod de pozare", "montaj"),
    "In_A": ("in impus", "in [a]", "in", "calibru", "protectie", "siguranta", "i [a]", "a"),
    "curba": ("curba", "curbă", "caracteristica"),
    "obs": ("observatii", "obs", "note", "mentiuni"),
}
SINONIME_RECEPTOR = {
    "iluminat": ("iluminat", "lumina", "lumini", "corpuri de iluminat", "il", "spoturi"),
    "prize": ("prize", "priza", "prize camere", "prize living", "prize dormitor", "prize generale"),
    "prize_bucatarie": ("prize bucatarie", "bucatarie"),
    "prize_baie": ("prize baie", "baie"),
    "prize_exterior": ("prize exterior", "exterior", "prize curte", "terasa"),
    "masina_spalat_rufe": ("masina de spalat rufe", "masina spalat rufe", "msr", "spalat rufe"),
    "masina_spalat_vase": ("masina de spalat vase", "masina spalat vase", "msv", "spalat vase"),
    "uscator": ("uscator", "uscator rufe", "uscator de rufe"),
    "boiler": ("boiler", "boiler electric", "apa calda"),
    "cuptor": ("cuptor", "cuptor electric"),
    "plita": ("plita", "plita electrica", "plita inductie", "aragaz electric"),
    "aer_conditionat": ("aer conditionat", "ac", "clima", "climatizare", "split"),
    "pompa_caldura": ("pompa de caldura", "pompa caldura", "pc", "heat pump"),
    "centrala_termica": ("centrala termica", "centrala", "ct"),
    "incalzire_electrica": ("incalzire electrica", "calorifer electric", "convector", "incalzire pardoseala"),
    "vehicul_electric": ("vehicul electric", "ve", "statie incarcare", "incarcator ve", "wallbox", "masina electrica"),
    "pompa_apa": ("pompa apa", "hidrofor", "pompa", "pompa submersibila"),
    "poarta_automata": ("poarta automata", "poarta", "automatizare poarta"),
    "sauna": ("sauna",),
    "tablou_secundar": ("tablou secundar", "td", "tablou"),
    "alt_circuit": ("alt circuit", "altele", "diverse", "rezerva"),
}
EXEMPLE = [("Iluminat parter", "Iluminat", 0.8, 0.95, 25, "L1", "B2", None, "B", ""),
           ("Prize bucătărie", "Prize bucătărie", 3.5, 0.95, 15, "L2", "B2", None, "", "blat"),
           ("Pompă de căldură", "Pompă de căldură", 6, 0.9, 20, "3F", "C", None, "", "unitatea exterioară")]


def _fd(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s.strip().lower())


def model_xlsx():
    """modelul de fișier de import (antet + 3 rânduri exemplu + listele permise)"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Circuite"
    ws.append([c[1] for c in COLOANE])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E6DA")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for e in EXEMPLE:
        ws.append(list(e))
    for i, (_, _, w) in enumerate(COLOANE, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wr = wb.create_sheet("Receptori")
    wr.append(["Tip receptor (de scris în coloana B)", "Cheie internă"])
    for k, v in ETICHETE_RECEPTOR.items():
        wr.append([v, k])
    wr.column_dimensions["A"].width = 36
    wr.column_dimensions["B"].width = 24
    n = len(ETICHETE_RECEPTOR) + 1
    dv = DataValidation(type="list", formula1=f"=Receptori!$A$2:$A${n}", allow_blank=True)
    dv.error, dv.errorTitle = "Alege un tip din foaia Receptori (sau lasă gol = „Alt circuit”).", "Tip receptor"
    ws.add_data_validation(dv)
    dv.add("B2:B500")
    dvf = DataValidation(type="list", formula1='"' + ",".join(FAZE) + '"', allow_blank=True)
    ws.add_data_validation(dvf)
    dvf.add("F2:F500")
    dvp = DataValidation(type="list", formula1='"' + ",".join(POZARI) + '"', allow_blank=True)
    ws.add_data_validation(dvp)
    dvp.add("G2:G500")
    wi = wb.create_sheet("Instrucțiuni")
    for t in ("Model de import al circuitelor — Scheme monofilare",
              "",
              "1. Un rând = un circuit. Rândurile exemplu (2–4) se șterg sau se înlocuiesc.",
              "2. Obligatoriu: Circuit (denumire). Restul coloanelor sunt opționale: goale = valorile uzuale ale aplicației "
              "(marcate „presupus” pe planșă și în memoriu).",
              "3. Tip receptor: din lista foii Receptori; determină curba, secțiunea minimă, grupa DDR și puterea implicită.",
              "4. P [kW] — puterea instalată a circuitului; se acceptă și W („2500 W”).",
              "5. Fază: L1 / L2 / L3 la monofazat (sau gol = repartizare automată), 3F = circuit trifazat.",
              "6. Pozare: B2 (în tub pe / în perete), C (aparent), D1 (în tub în pământ), D2 (direct în pământ), E (pe pat).",
              "7. In impus [A]: doar dacă vreți un calibru anume; aplicația îl verifică (Iz, ΔU, deconectare).",
              "8. Se poate importa și ca CSV (separator „;”), cu același antet."):
        wi.append([t])
    wi.column_dimensions["A"].width = 120
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _num(v):
    """„2,5”, „16 A”, „3kW”, „2500 W” → float (W → kW), None pentru gol / text"""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(" ", " ")
    m = re.search(r"(-?\d+(?:[.,]\d+)?)\s*(kw|w|a|m|ma)?\b", s, re.I)
    if not m:
        return None
    x = float(m.group(1).replace(",", "."))
    if (m.group(2) or "").lower() == "w" and x >= 50:
        x = x / 1000.0
    return x


def _receptor(v):
    t = _fd(v)
    if not t:
        return None
    if t.replace(" ", "_") in ETICHETE_RECEPTOR:
        return t.replace(" ", "_")
    for k, et in ETICHETE_RECEPTOR.items():
        if _fd(et) == t:
            return k
    for k, sin in SINONIME_RECEPTOR.items():
        if t in sin:
            return k
    for k, sin in SINONIME_RECEPTOR.items():      # conține un sinonim lung (≥ 4 litere)
        if any(len(s) >= 4 and re.search(rf"\b{re.escape(s)}\b", t) for s in sin):
            return k
    return False


def _randuri(date):
    """bytes (xlsx sau csv) sau text → listă de rânduri (liste de valori)"""
    if isinstance(date, str):
        date = date.encode("utf-8")
    if date[:2] == b"PK":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(date), data_only=True, read_only=True)
        ws = wb["Circuite"] if "Circuite" in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if date[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError("fișierul e .xls (Excel 97–2003) — salvați-l ca .xlsx sau .csv")
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = date.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    linii = [x for x in text.splitlines() if x.strip()]
    if not linii:
        return []
    prima = linii[0]
    sep = max((";", "\t", ","), key=lambda c: prima.count(c))
    return [r for r in csv.reader(io.StringIO("\n".join(linii)), delimiter=sep)]


def _antet(rows):
    """(indexul rândului de antet, {cheie: coloană})"""
    best = (None, {})
    for i, r in enumerate(rows[:10]):
        harta = {}
        for j, cel in enumerate(r):
            t = _fd(cel)
            if not t:
                continue
            for k, sin in SINONIME_COLOANE.items():
                if k in harta:
                    continue
                if any(t == s or t.startswith(s + " ") or t.startswith(s + "[") or t.startswith(s + "(") for s in sin):
                    harta[k] = j
                    break
        if "nume" in harta and len(harta) > len(best[1]):
            best = (i, harta)
    return best


def importa(date, destinatie=None):
    """date = conținutul fișierului; destinatie = „locuinta” (circuite) / „tablou” (plecări) / None (ambele)"""
    rows = _randuri(date)
    if not rows:
        raise ValueError("fișierul nu conține date")
    ih, harta = _antet(rows)
    if ih is None:
        raise ValueError("nu am găsit antetul (o coloană „Circuit” / „Denumire”) în primele 10 rânduri — folosiți modelul .xlsx")
    av, ign, circ, plec = [], [], [], []
    for nr, r in enumerate(rows[ih + 1:], ih + 2):
        def v(k):
            j = harta.get(k)
            return r[j] if j is not None and j < len(r) else None
        nume = str(v("nume") or "").strip()
        if not nume or re.match(r"^(total|subtotal|suma)\b", _fd(nume)):
            if any(x not in (None, "") for x in r):
                ign.append({"rand": nr, "motiv": "fără denumire / rând de total"})
            continue
        c = {"nume": nume}
        rec = _receptor(v("receptor")) if v("receptor") not in (None, "") else _receptor(nume) or None
        if rec is False:
            av.append(f"rândul {nr} ({nume}): tipul „{v('receptor')}” nu e recunoscut → „Alt circuit”")
            rec = "alt_circuit"
        if rec:
            c["receptor"] = rec
        for k in ("P_kW", "cosfi", "L_m", "In_A"):
            x = _num(v(k))
            if x is not None:
                if x <= 0 or (k == "cosfi" and x > 1):
                    av.append(f"rândul {nr} ({nume}): {k} = {v(k)} ignorat (valoare nevalidă)")
                    continue
                c[k] = int(x) if float(x).is_integer() else round(x, 4)
        f = str(v("faza") or "").strip().upper().replace(" ", "")
        if f:
            f = {"3": "3F", "3F": "3F", "TRI": "3F", "TRIFAZAT": "3F", "R": "L1", "S": "L2", "T": "L3"}.get(f, f)
            if f in FAZE:
                c["faza"] = f
            else:
                av.append(f"rândul {nr} ({nume}): faza „{v('faza')}” nu e L1/L2/L3/3F — repartizare automată")
        met = str(v("metoda") or "").strip().upper()
        if met:
            if met in POZARI:
                c["cablu"] = {"metoda": met}
            else:
                av.append(f"rândul {nr} ({nume}): pozarea „{v('metoda')}” necunoscută — implicit")
        cb = str(v("curba") or "").strip().upper()
        if cb in ("B", "C", "D"):
            c["curba"] = cb
        circ.append(c)
        p = {"nume": nume}
        for k in ("P_kW", "In_A", "L_m", "cosfi"):
            if k in c:
                p[k] = c[k]
        if c.get("faza") == "3F" or (c.get("faza") is None and (c.get("P_kW") or 0) > 7.4):
            p["retea"] = "tri"
        elif c.get("faza"):
            p["retea"], p["faza"] = "mono", c["faza"]
        if "cablu" in c:
            p["cablu"] = dict(c["cablu"])
        if not (p.get("P_kW") or p.get("In_A")):
            av.append(f"rândul {nr} ({nume}): fără P sau In — la tablou plecarea cere una dintre ele")
        plec.append(p)
    out = {"n": len(circ), "avertismente": av, "ignorate": ign, "coloane_recunoscute": sorted(harta)}
    if destinatie in (None, "locuinta"):
        out["circuite"] = circ
    if destinatie in (None, "tablou", "cascada"):
        out["plecari"] = plec
    return out
