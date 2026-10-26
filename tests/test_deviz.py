"""Deviz estimativ, catalog de prețuri, pachet de predare, profile."""
import io, json, os, pathlib, sys, tempfile, threading, urllib.request, urllib.error, zipfile
import pytest

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
_TMP = tempfile.mkdtemp()
os.environ["SM_DATE"] = _TMP          # testele nu ating catalogul real al utilizatorului
import importlib
from engine import deviz as _d
importlib.reload(_d)
from engine import constructor, layout, deviz  # noqa: E402

INTRARE = {"tip_schema": "tablou", "P_kW": 120, "L_alimentare_m": 40,
           "plecari": [{"nume": "TD-2", "In_A": 63}, {"nume": "Prize", "In_A": 16}]}


def _poz():
    m, _ = constructor.genereaza(INTRARE)
    return m, layout.build_cu_pozitii(m)[1]


def test_catalog_in_folder_temporar():
    assert str(deviz.CATALOG).startswith(_TMP)


def test_fara_preturi_nu_inventeaza():
    m, poz = _poz()
    d = deviz.calculeaza(poz, catalog={"articole": {}, "setari": dict(deviz.SETARI_IMPLICITE)})
    assert d["recapitulatie"]["materiale"] == 0 and d["recapitulatie"]["fara_pret"] == len(poz)
    assert all(a["pret"] is None and a["valoare"] is None for a in d["articole"])


def test_recapitulatie_aritmetica():
    m, poz = _poz()
    art = deviz.articole(poz)
    cat = {"articole": {a["cheie"]: {"pret": 10.0, "ore": 1.0} for a in art},
           "setari": dict(deviz.SETARI_IMPLICITE, tarif_ora=50, ch_indirecte_pct=10, profit_pct=5, tva_pct=21)}
    d = deviz.calculeaza(poz, catalog=cat)
    q = sum(a["cantitate"] for a in art)
    r = d["recapitulatie"]
    assert r["materiale"] == pytest.approx(10 * q) and r["ore"] == pytest.approx(q)
    assert r["manopera"] == pytest.approx(50 * q)
    baza = 60 * q
    assert r["ch_indirecte"] == pytest.approx(baza * 0.10, abs=0.01)
    assert r["profit"] == pytest.approx(baza * 1.10 * 0.05, abs=0.01)
    assert r["total"] == pytest.approx(r["total_fara_tva"] * 1.21, abs=0.02)


def test_cheie_ignora_spatii_si_majuscule():
    assert deviz.cheie("Cablu  ACYABY 4×25 mm² Al") == deviz.cheie("cablu acyaby 4×25 mm² al")


def test_catalog_scriere_si_import_xlsx():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Cod", "Denumire articol", "Preț unitar", "Furnizor"])
    ws.append(["X1", "Cablu ACYABY 4×25 mm² Al", "12,5", "Elcond"])
    ws.append(["", "", "", ""])
    bio = io.BytesIO(); wb.save(bio)
    nou = deviz.importa_xlsx(bio.getvalue())
    assert nou["articole"][deviz.cheie("Cablu ACYABY 4×25 mm² Al")]["pret"] == 12.5
    d = deviz.scrie_catalog(nou)
    assert deviz.citeste_catalog()["articole"][deviz.cheie("cablu acyaby 4×25 mm² al")]["furnizor"] == "Elcond"
    deviz.scrie_catalog({"articole": {"Cablu ACYABY 4×25 mm² Al": None}})
    assert deviz.cheie("Cablu ACYABY 4×25 mm² Al") not in deviz.citeste_catalog()["articole"]


def test_deviz_xlsx_se_deschide():
    from openpyxl import load_workbook
    m, poz = _poz()
    d = deviz.calculeaza(poz, catalog={"articole": {}, "setari": dict(deviz.SETARI_IMPLICITE)})
    wb = load_workbook(io.BytesIO(deviz.deviz_xlsx(m, d)))
    assert wb.sheetnames == ["Deviz", "Recapitulație"]
    assert wb["Deviz"].max_row >= 6 + len(poz)


def test_lista_contine_aparataj_auxiliar():
    m, poz = _poz()
    den = [p[2][0] for p in poz]
    assert any(x.startswith("Patron fuzibil") for x in den)
    assert any(x.startswith("Sistem de bare") for x in den)
    assert any(x.startswith("Carcasă tablou") for x in den)


@pytest.fixture(scope="module")
def server():
    import app
    importlib.reload(app)
    from http.server import ThreadingHTTPServer
    srv = ThreadingHTTPServer(("127.0.0.1", 0), app.H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()


def _post(url, body, raw=False):
    data = body if raw else json.dumps(body).encode()
    req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/octet-stream" if raw else "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def test_api_pachet_si_deviz(server):
    st, b = _post(server + "/api/genereaza", {"calc": INTRARE})
    m = json.loads(b)["model"]
    m["deviz"] = {"setari": {"tarif_ora": 40}, "suprascrieri": {}}
    st, b = _post(server + "/api/deviz", m)
    d = json.loads(b)
    assert st == 200 and d["setari"]["tarif_ora"] == 40 and d["recapitulatie"]["nr_articole"] > 5
    st, b = _post(server + "/api/export/pachet", m)
    assert st == 200
    z = zipfile.ZipFile(io.BytesIO(b))
    nume = z.namelist()
    for suf in (".dxf", ".svg", "_lista_echipamente.xlsx", "_deviz_estimativ.xlsx", "_proiect.json", "CITESTE.txt"):
        assert any(n.endswith(suf) for n in nume), suf
    assert any("memoriu" in n for n in nume)


def test_api_pachet_blocat(server):
    st, b = _post(server + "/api/genereaza", {"calc": {"tip_schema": "bransament", "retea": "tri", "P_kW": 300, "L_coloana_m": 80}})
    st, b = _post(server + "/api/export/pachet", json.loads(b)["model"])
    assert st == 409


def test_api_profile(server):
    st, b = _post(server + "/api/profile", {"profile": [{"nume": "Firma A", "proiectant": "A SRL"}, {"proiectant": "fără nume"}], "activ": "Firma A"})
    d = json.loads(b)
    assert [p["nume"] for p in d["profile"]] == ["Firma A"] and d["activ"] == "Firma A"
    with urllib.request.urlopen(server + "/api/profile") as r:
        assert json.loads(r.read())["profile"][0]["proiectant"] == "A SRL"


def _xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    for r in rows:
        ws.append(r)
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


@pytest.mark.parametrize("antet,valoare,asteptat", [
    (["Denumire", "Cantitate", "Preț total", "Preț unitar"], [None, 10, 1115, 111.5], 111.5),        # C5-1: nu „total”
    (["Denumire", "Preț cu TVA", "Preț fără TVA"], [None, 134.92, 111.5], 111.5),                   # C5-1: nu „cu TVA”
    (["Denumire", "Preț unitar"], [None, "1.250"], 1250.0),                                          # C5-2: mii
    (["Denumire", "Preț unitar"], [None, "1.250,50"], 1250.5),
    (["Denumire", "Preț unitar"], [None, "111,5 lei"], 111.5),
    (["Articol", "PU"], [None, 12], None),
])
def test_import_preturi_format_romanesc(antet, valoare, asteptat):
    row = list(valoare); row[0] = "Cablu ACYABY 4×240 mm² Al"
    if asteptat is None:
        with pytest.raises(ValueError):
            deviz.importa_xlsx(_xlsx([antet, row]))
        return
    r = deviz.importa_xlsx(_xlsx([antet, row]))
    assert r["articole"][deviz.cheie(row[0])]["pret"] == asteptat


def test_import_doar_cu_tva_refuzat():
    with pytest.raises(ValueError):
        deviz.importa_xlsx(_xlsx([["Denumire", "Preț cu TVA"], ["Cablu X", 10]]))


def test_setari_invalide_refuzate():
    for st in ({"tarif_ora": "abc"}, {"tva_pct": -21}, {"profit_pct": 150}):
        with pytest.raises(ValueError):
            deviz.scrie_catalog({"setari": st})
    with pytest.raises(ValueError):
        deviz.scrie_catalog({"articole": {"Cablu X": {"pret": -50}}})
    m, poz = _poz()
    with pytest.raises(ValueError):
        deviz.calculeaza(poz, catalog={"articole": {}, "setari": dict(deviz.SETARI_IMPLICITE)}, setari={"tarif_ora": "x"})
